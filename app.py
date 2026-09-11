import io
import re
from datetime import datetime
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONTH_NAMES = {
    1: "JANUARY", 2: "FEBRUARY", 3: "MARCH", 4: "APRIL",
    5: "MAY", 6: "JUNE", 7: "JULY", 8: "AUGUST",
    9: "SEPTEMBER", 10: "OCTOBER", 11: "NOVEMBER", 12: "DECEMBER"
}

def clean_int(val):
    """અલ્પવિરામ કે સ્પેસ દૂર કરીને સાચો પૂર્ણાંક નંબર બનાવે છે."""
    s = re.sub(r"[^\d]", "", str(val))
    return int(s) if s else 0

def normalize_date(date_str):
    clean_d = str(date_str).strip().replace("/", "-")
    parts = clean_d.split("-")
    if len(parts) == 3:
        try:
            d = int(parts[0])
            m = int(parts[1])
            y = int(parts[2])
            if y < 100:
                y += 2000
            return f"{str(d).zfill(2)}-{str(m).zfill(2)}-{y}"
        except Exception:
            pass
    return clean_d

def parse_date_obj(date_str):
    clean_d = normalize_date(date_str)
    parts = clean_d.split("-")
    if len(parts) == 3:
        try:
            return datetime(int(parts[2]), int(parts[1]), int(parts[0]))
        except Exception:
            pass
    return datetime(2099, 1, 1)

def parse_raw_text(text):
    pattern = re.compile(
        r"PROJECT\s*(?:NAME)?\s*[:=-]+\s*(?P<project>[^\n\r]+)[\r\n]+"
        r"EMPLOYEE\s*(?:NAME)?\s*[:=-]+\s*(?P<employee>[^\n\r]+)[\r\n]+"
        r"DATE\s*[:=-]+\s*(?P<date>[^\n\r]+)[\r\n]+"
        r"(?:SCAN\s*)?FILE\s*[:=-]+\s*(?P<files>[\d,\s]+)[\r\n]+"
        r"(?:SCAN\s*)?PAGE\s*[:=-]+\s*(?P<pages>[\d,\s]+)",
        re.IGNORECASE,
    )

    entries = []
    for m in pattern.finditer(text):
        entries.append({
            "Project": m.group("project").strip().upper(),
            "Employee": m.group("employee").strip().upper(),
            "Date": normalize_date(m.group("date")),
            "Scan File": clean_int(m.group("files")),
            "Scan Page": clean_int(m.group("pages")),
        })

    if not entries:
        alt_pattern = re.compile(
            r"PROJECT\s*(?:NAME)?\s*[:=-]+\s*(?P<project>[^\n\r]+)[\r\n]+"
            r"(?:EMPLOYEE\s*(?:NAME)?\s*[:=-]+\s*(?P<employee>[^\n\r]+)[\r\n]+)?"
            r"DATE\s*[:=-]+\s*(?P<date>[^\n\r]+)[\r\n]+"
            r"(?:SCAN\s*)?PAGE\s*[:=-]+\s*(?P<pages>[\d,\s]+)[\r\n]+"
            r"(?:SCAN\s*)?FILE\s*[:=-]+\s*(?P<files>[\d,\s]+)",
            re.IGNORECASE,
        )
        for m in alt_pattern.finditer(text):
            emp = m.group("employee")
            emp = emp.strip().upper() if emp else "UNKNOWN"
            entries.append({
                "Project": m.group("project").strip().upper(),
                "Employee": emp,
                "Date": normalize_date(m.group("date")),
                "Scan File": clean_int(m.group("files")),
                "Scan Page": clean_int(m.group("pages")),
            })

    return entries

def generate_project_filename(records):
    if not records:
        return "SCANNING_REPORT.xlsx"

    project_name = records[0]["Project"].replace(" ", "_")
    parsed_dates = [parse_date_obj(r["Date"]) for r in records if parse_date_obj(r["Date"]).year != 2099]

    if not parsed_dates:
        return f"{project_name}_REPORT.xlsx"

    parsed_dates.sort()
    months_in_order = []
    for d in parsed_dates:
        m_name = MONTH_NAMES.get(d.month, str(d.month))
        if m_name not in months_in_order:
            months_in_order.append(m_name)

    years = sorted(list(set(d.year for d in parsed_dates)))
    year_str = "-".join(str(y) for y in years)
    month_str = months_in_order[0] if len(months_in_order) == 1 else "-".join(months_in_order)
    return f"{project_name}_{month_str}_{year_str}.xlsx"
def generate_excel_bytes(records):
    df = pd.DataFrame(records)
    wb = Workbook()

    ws_master = wb.active
    ws_master.title = "Master Data"
    ws_master.append(["Project", "Employee", "Date", "Scan File", "Scan Page"])
    for r in records:
        ws_master.append([r["Project"], r["Employee"], r["Date"], r["Scan File"], r["Scan Page"]])

    ws = wb.create_sheet(title="Scanning Sheet", index=0)
    ws.views.sheetView[0].showGridLines = True

    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF")
    )
    thick_bottom = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="medium", color="000000")
    )

    font_name_col = Font(name="Calibri", size=11, bold=True, color="C00000")
    font_date = Font(name="Calibri", size=11, bold=True, color="C00000")
    font_emp = Font(name="Calibri", size=11, bold=True, color="1F4E78")
    font_file_label = Font(name="Calibri", size=10, bold=True, color="333333")
    font_data = Font(name="Calibri", size=11, bold=False, color="000000")
    font_total = Font(name="Calibri", size=11, bold=True, color="002060")

    fill_total = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    unique_dates = sorted(list(set(df["Date"].tolist())), key=parse_date_obj)
    employees = sorted(df["Employee"].unique().tolist())

    c_name = ws.cell(row=2, column=1, value="NAME")
    c_name.font = font_name_col
    c_name.alignment = Alignment(horizontal="center", vertical="center")
    c_name.border = thin_border
    ws.cell(row=2, column=2, value="").border = thin_border

    date_col_map = {}
    col_idx = 3
    for d_str in unique_dates:
        d_obj = parse_date_obj(d_str)
        disp_date = f"{d_obj.day}/{d_obj.month}/{d_obj.year}" if d_obj.year != 2099 else d_str
        c = ws.cell(row=2, column=col_idx, value=disp_date)
        c.font = font_date
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        date_col_map[d_str] = col_idx
        col_idx += 1

    total_col_idx = col_idx
    c_tot = ws.cell(row=2, column=total_col_idx, value="TOTAL")
    c_tot.font = font_name_col
    c_tot.alignment = Alignment(horizontal="center", vertical="center")
    c_tot.border = thin_border
    c_tot.fill = fill_total
    ws.row_dimensions[2].height = 25

    data_lookup = {(r["Employee"], r["Date"]): (r["Scan File"], r["Scan Page"]) for _, r in df.iterrows()}

    current_row = 4
    for emp in employees:
        page_row = current_row
        file_row = current_row + 1

        c_emp = ws.cell(row=page_row, column=1, value=emp)
        c_emp.font = font_emp
        c_emp.alignment = Alignment(horizontal="left", vertical="center")
        c_emp.border = thin_border
        ws.cell(row=page_row, column=2, value="").border = thin_border

        c_file_lbl = ws.cell(row=file_row, column=1, value="FILE")
        c_file_lbl.font = font_file_label
        c_file_lbl.alignment = Alignment(horizontal="center", vertical="center")
        c_file_lbl.border = thick_bottom
        ws.cell(row=file_row, column=2, value="").border = thick_bottom

        for d_str, c_num in date_col_map.items():
            file_val, page_val = data_lookup.get((emp, d_str), ("", ""))

            cp = ws.cell(row=page_row, column=c_num, value=page_val if page_val != "" else "")
            cp.font = font_data
            cp.alignment = Alignment(horizontal="center", vertical="center")
            cp.border = thin_border

            cf = ws.cell(row=file_row, column=c_num, value=file_val if file_val != "" else "")
            cf.font = font_data
            cf.alignment = Alignment(horizontal="center", vertical="center")
            cf.border = thick_bottom

        first_col_let = get_column_letter(3)
        last_col_let = get_column_letter(total_col_idx - 1)

        tot_p = ws.cell(row=page_row, column=total_col_idx, value=f"=SUM({first_col_let}{page_row}:{last_col_let}{page_row})")
        tot_p.font = font_total
        tot_p.alignment = Alignment(horizontal="center", vertical="center")
        tot_p.border = thin_border
        tot_p.fill = fill_total

        tot_f = ws.cell(row=file_row, column=total_col_idx, value=f"=SUM({first_col_let}{file_row}:{last_col_let}{file_row})")
        tot_f.font = font_total
        tot_f.alignment = Alignment(horizontal="center", vertical="center")
        tot_f.border = thick_bottom
        tot_f.fill = fill_total

        ws.row_dimensions[page_row].height = 20
        ws.row_dimensions[file_row].height = 20
        current_row += 2

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 5
    for c_i in range(3, total_col_idx + 1):
        col_let = get_column_letter(c_i)
        ws.column_dimensions[col_let].width = 13

    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream

# ==================== Streamlit Web App ====================
st.set_page_config(page_title="Ajay Valiya Excel Data Entry", layout="wide", page_icon="📊")

st.markdown("""
    <h2 style='text-align: center; color: #1F4E78;'>AJAY VALIYA EXCEL DATA ENTRY SYSTEM</h2>
    <p style='text-align: center; color: gray;'>વોટ્સએપ મેસેજ પેસ્ટ કરો અને તારીખ વાઈઝ સુંદર એક્સેલ બનાવો</p>
""", unsafe_allow_html=True)

if "master_records" not in st.session_state:
    st.session_state.master_records = {}

if "raw_text_input" not in st.session_state:
    st.session_state.raw_text_input = ""

st.sidebar.header("📁 અગાઉ બનાવેલી Excel અપલોડ કરો")
uploaded_file = st.sidebar.file_uploader("જો જૂનો ડેટા જોડવો હોય તો ફાઈલ ચૂઝ કરો:", type=["xlsx"])

if uploaded_file and "file_loaded" not in st.session_state:
    try:
        df_old = pd.read_excel(uploaded_file, sheet_name="Master Data")
        for _, r in df_old.iterrows():
            d_norm = normalize_date(r["Date"])
            key = (str(r["Project"]).strip().upper(), str(r["Employee"]).strip().upper(), d_norm)
            st.session_state.master_records[key] = {
                "Project": str(r["Project"]).strip().upper(),
                "Employee": str(r["Employee"]).strip().upper(),
                "Date": d_norm,
                "Scan File": clean_int(r["Scan File"]),
                "Scan Page": clean_int(r["Scan Page"])
            }
        st.session_state.file_loaded = True
        st.sidebar.success(f"{len(df_old)} જૂના રેકોર્ડ્સ લોડ થયા!")
    except Exception:
        st.sidebar.error("Master Data શીટ વાંચવામાં ભૂલ આવી.")

raw_input = st.text_area(
    "વોટ્સએપ રો ડેટા અહીં પેસ્ટ કરો:",
    value=st.session_state.raw_text_input,
    key="raw_text_area",
    height=250,
    placeholder="PROJECT NAME:- NARODA\nEMPLOYEE NAME:- Kinjal\nDATE:- 01-09-2026\nSCAN FILE:- 24\nSCAN PAGE :- 2,027..."
)

col1, col2 = st.columns([2, 1])

with col1:
    process_btn = st.button("🚀 ડેટા પ્રોસેસ કરો", type="primary", use_container_width=True)

with col2:
    if st.button("🧹 ટેક્સ્ટ ક્લિયર કરો", use_container_width=True):
        st.session_state.raw_text_input = ""
        st.rerun()

if process_btn and raw_input.strip():
    new_entries = parse_raw_text(raw_input)
    if not new_entries:
        st.error("ટેક્સ્ટમાંથી કોઈ યોગ્ય ડેટા મળ્યો નહીં. ફોર્મેટ ચેક કરો.")
    else:
        added = 0
        skipped = 0
        conflicts = []

        for entry in new_entries:
            key = (entry["Project"], entry["Employee"], entry["Date"])
            if key in st.session_state.master_records:
                old = st.session_state.master_records[key]
                if old["Scan File"] == entry["Scan File"] and old["Scan Page"] == entry["Scan Page"]:
                    skipped += 1
                else:
                    conflicts.append((entry, old))
            else:
                st.session_state.master_records[key] = entry
                added += 1

        st.session_state.temp_conflicts = conflicts
        st.success(f"નવી એન્ટ્રી ઉમેરાઈ: {added} | સમાન ડેટા સ્કીપ થયો: {skipped}")

if "temp_conflicts" in st.session_state and st.session_state.temp_conflicts:
    st.warning("⚠️ નીચેની એન્ટ્રીઓમાં તારીખ સેમ છે પરંતુ File/Page ના આંકડા બદલાયેલા છે:")
    for idx, (new_d, old_d) in enumerate(st.session_state.temp_conflicts):
        st.write(f"**પ્રોજેક્ટ**: {new_d['Project']} | **કર્મચારી**: {new_d['Employee']} | **તારીખ**: {new_d['Date']}")
        st.write(f"• **જૂનો ડેટા**: Files={old_d['Scan File']}, Pages={old_d['Scan Page']}")
        st.write(f"• **નવો ડેટા**: Files={new_d['Scan File']}, Pages={new_d['Scan Page']}")
        
        choice = st.radio(f"કયો ડેટા રાખવો છે? (#{idx+1})", ["નવો ડેટા રાખો (Replace)", "જૂનો ડેટા રાખો (Keep Old)"], key=f"conf_{idx}")
        if choice == "નવો ડેટા રાખો (Replace)":
            k = (new_d["Project"], new_d["Employee"], new_d["Date"])
            st.session_state.master_records[k] = new_d

if st.session_state.master_records:
    records_list = list(st.session_state.master_records.values())
    file_name = generate_project_filename(records_list)
    excel_data = generate_excel_bytes(records_list)

    st.markdown("---")
    st.subheader("📥 તૈયાર Excel ફાઇલ ડાઉનલોડ કરો:")
    st.download_button(
        label=f"⬇️ {file_name} ડાઉનલોડ કરો",
        data=excel_data,
        file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
